"""Unit tests for the multi-format parsers and segment chunking (ADR-0018). No services needed."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from alexsoft_rag.errors import ApiError
from alexsoft_rag.ingest.chunker import chunk_segments
from alexsoft_rag.parsers import parse_document, supported_extensions
from alexsoft_rag.parsers.base import ParseContext, RawDocument
from alexsoft_rag.schemas import ChunkingConfig


async def parse(name: str, data: bytes | str, **ctx) -> list:
    raw = RawDocument(name, data.encode("utf-8") if isinstance(data, str) else data)
    return await parse_document(raw, ParseContext(**ctx))


def test_registry_covers_new_formats():
    exts = set(supported_extensions())
    assert {".py", ".ts", ".tsx", ".ps1", ".sql", ".yaml", ".dsl", ".puml", ".pdf", ".xlsx", ".bpmn", ".drawio"} <= exts


# ---- code ---------------------------------------------------------------------------
PY = '''"""Module doc."""
import os


def alpha(x):
    """First function."""
    return x + 1


class Service:
    """A service."""

    def embed_query(self, text):
        return text.lower()

    def other(self):
        return 2
'''


async def test_python_symbols_and_line_numbers():
    segs = await parse("apps/rag/svc.py", PY)
    assert all(s.kind == "code" for s in segs)
    joined = "\n".join(s.text for s in segs)
    assert "def alpha" in joined and "class Service" in joined
    first = segs[0]
    assert first.header.startswith("File: apps/rag/svc.py\nLanguage: python")
    assert first.locator["line_start"] == 1
    lines = PY.split("\n")
    for s in segs:  # locator lines point at the original text
        a, b = s.locator["line_start"], s.locator["line_end"]
        assert "\n".join(lines[a - 1 : b]) == s.text


async def test_python_large_class_is_split_into_methods():
    body = "\n".join(f"    def method_{i}(self):\n        return {i}  # " + "x" * 80 for i in range(30))
    src = f"class Big:\n    \"\"\"doc\"\"\"\n\n{body}\n"
    segs = await parse("big.py", src, chunk_size=128)
    symbols = [s.locator.get("symbol", "") for s in segs]
    assert any("Big.method_" in s for s in symbols)
    assert len(segs) > 3


async def test_typescript_and_tsx():
    ts = "export interface A { x: number }\n\nexport function runJob(id: string) {\n  return id;\n}\n\nexport const useThing = () => 1;\n"
    segs = await parse("a.ts", ts)
    assert segs and "runJob" in segs[0].locator.get("symbol", "") + segs[0].text
    tsx = "export default function Page() {\n  return <div>hi</div>;\n}\n"
    segs = await parse("page.tsx", tsx)
    assert "Page" in segs[0].locator["symbol"]


async def test_typescript_large_class_is_split_into_methods():
    methods = "\n".join(
        f"  async handle{i}(req: Request): Promise<void> {{\n    await this.svc.run({i}); // " + "y" * 60 + "\n  }\n"
        for i in range(25)
    )
    src = f"import {{ Request }} from './x';\n\n/** The controller. */\nexport class Controller {{\n{methods}}}\n"
    segs = await parse("ctl.ts", src, chunk_size=128)
    symbols = [s.locator.get("symbol", "") for s in segs]
    assert any("Controller.handle3" in s for s in symbols)
    lines = src.rstrip("\n").split("\n")
    for s in segs:
        assert "\n".join(lines[s.locator["line_start"] - 1 : s.locator["line_end"]]) == s.text


async def test_unparsable_python_falls_back_to_blocks():
    segs = await parse("broken.py", "def ok():\n    return 1\n\ndef broken(:\n    pass\n")
    assert segs and "ok" in segs[0].locator["symbol"]


async def test_block_languages_have_symbols():
    ps1 = "param()\n\nfunction Start-Rag {\n  Write-Host 'x'\n}\n\nfunction Stop-Rag {\n  Write-Host 'y'\n}\n"
    segs = await parse("start.ps1", ps1, chunk_size=16)
    symbols = ", ".join(s.locator.get("symbol", "") for s in segs)
    assert "Start-Rag" in symbols and "Stop-Rag" in symbols
    # big blocks stay separate units, each with its own symbol
    fat = "\n".join(f"  Write-Host 'line {i} with some padding text to take tokens'" for i in range(40))
    segs = await parse("fat.ps1", f"function A {{\n{fat}\n}}\n\nfunction B {{\n{fat}\n}}\n", chunk_size=128)
    assert [s.locator["symbol"] for s in segs] == ["A", "B"]
    yml = "services:\n  a: 1\nvolumes:\n  b: 2\n"
    segs = await parse("compose.yaml", yml)
    assert "services" in segs[0].locator["symbol"] and "volumes" in segs[0].locator["symbol"]
    sql = "CREATE TABLE rag.docs (id int);\n\nCREATE INDEX idx_docs ON rag.docs (id);\n"
    segs = await parse("m.sql", sql)
    assert "rag.docs" in segs[0].locator["symbol"] and "idx_docs" in segs[0].locator["symbol"]


async def test_oversized_code_unit_is_packed_with_correct_lines():
    src = "def f():\n" + "\n".join(f"    v{i} = {i}" for i in range(400)) + "\n"
    segs = await parse("f.py", src, chunk_size=512)
    chunks = chunk_segments(segs, ChunkingConfig(chunk_size=128, chunk_overlap=0))
    assert len(chunks) > 3
    lines = src.rstrip("\n").split("\n")
    for c in chunks:
        a, b = c.locator["line_start"], c.locator["line_end"]
        assert c.text == "\n".join(lines[a - 1 : b])
        assert f"Lines: {a}-{b}" in c.header


async def test_non_utf8_code_is_rejected():
    with pytest.raises(ApiError) as e:
        await parse("x.py", b"\xff\xfe\x00bad")
    assert e.value.status == 422


# ---- xlsx ------------------------------------------------------------------------------
def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Бюджет"
    ws.append(["Статья", "Q1", "Q2", "Итого"])
    for i in range(1, 41):
        ws.append([f"Статья {i}", i * 10, i * 20, f"=B{i + 1}+C{i + 1}"])
    ws.append(["Всего", "=SUM(B2:B41)", "=SUM(C2:C41)", "=SUM(D2:D41)"])
    hidden = wb.create_sheet("Скрытый")
    hidden.append(["секрет", 1])
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def test_xlsx_tables_and_formulas():
    segs = await parse("budget.xlsx", make_xlsx(), chunk_size=128)
    tables = [s for s in segs if "формулы" not in s.header]
    formulas = [s for s in segs if "формулы" in s.header]
    assert len(tables) > 1  # row groups
    assert all(s.text.startswith("| Статья | Q1 | Q2 | Итого |") for s in tables)  # header repeated
    assert all(s.locator["sheet"] == "Бюджет" and s.locator["range"].startswith("A") for s in tables)
    assert not any("секрет" in s.text for s in segs)  # hidden sheets are skipped
    assert formulas and any("=SUM(B2:B41)" in s.text for s in formulas)
    assert formulas[0].locator["range"].startswith("D2")
    chunks = chunk_segments(segs, ChunkingConfig(chunk_size=128, chunk_overlap=0))
    assert {c.kind for c in chunks} == {"table"}


async def test_broken_xlsx_is_422():
    with pytest.raises(ApiError) as e:
        await parse("bad.xlsx", b"not a zip")
    assert e.value.status == 422


# ---- pdf --------------------------------------------------------------------------------
def make_text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def make_mixed_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_page(PdfReader(io.BytesIO(make_text_pdf("RabbitMQ is the broker for long running AI tasks"))).pages[0])
    writer.add_blank_page(width=612, height=792)  # "scanned" page: no text layer
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


class FakeOCR:
    def __init__(self) -> None:
        self.calls = 0
        self.last_png = b""

    async def read_page(self, png: bytes, prompt: str | None = None) -> str:
        self.calls += 1
        self.last_png = png
        return "Scanned page: Qdrant is the vector database of the platform."


async def test_pdf_text_layer_and_ocr_for_scanned_pages():
    ocr = FakeOCR()
    segs = await parse("doc.pdf", make_mixed_pdf(), ocr=ocr)
    assert [s.locator["page"] for s in segs] == [1, 2]
    assert "RabbitMQ" in segs[0].text and "Qdrant" in segs[1].text
    assert segs[1].header == "doc.pdf, стр. 2"
    assert ocr.calls == 1  # only the page without a text layer goes to the vision model
    assert ocr.last_png.startswith(b"\x89PNG")


async def test_pdf_without_ocr_skips_scans_and_fails_if_nothing_left():
    segs = await parse("doc.pdf", make_mixed_pdf(), ocr=None)
    assert [s.locator["page"] for s in segs] == [1]
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(ApiError) as e:
        await parse("scan.pdf", out.getvalue(), ocr=None)
    assert e.value.status == 422


async def test_ocr_page_limit():
    writer = PdfWriter()
    for _ in range(4):
        writer.add_blank_page(width=100, height=100)
    out = io.BytesIO()
    writer.write(out)
    ocr = FakeOCR()
    segs = await parse("scan.pdf", out.getvalue(), ocr=ocr, max_ocr_pages=2)
    assert ocr.calls == 2 and [s.locator["page"] for s in segs] == [1, 2]


async def test_broken_pdf_is_422():
    with pytest.raises(ApiError) as e:
        await parse("bad.pdf", b"%PDF-1.4 garbage")
    assert e.value.status == 422


# ---- diagrams ------------------------------------------------------------------------------
DRAWIO = """<mxfile><diagram name="Target"><mxGraphModel><root>
<mxCell id="0"/><mxCell id="1" parent="0"/>
<mxCell id="c" value="Agent Platform" vertex="1" parent="1"/>
<mxCell id="a" value="LangGraph" vertex="1" parent="c"/>
<mxCell id="b" value="&lt;b&gt;LiteLLM&lt;/b&gt;" vertex="1" parent="1"/>
<mxCell id="e" value="chat" edge="1" source="a" target="b" parent="1"/>
</root></mxGraphModel></diagram></mxfile>"""

BPMN = """<?xml version="1.0"?>
<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL">
<bpmn:process id="p1" name="Обработка заявки">
<bpmn:startEvent id="s" name="Заявка получена"/>
<bpmn:userTask id="t" name="Проверить заявку"/>
<bpmn:endEvent id="e" name="Готово"/>
<bpmn:sequenceFlow id="f1" sourceRef="s" targetRef="t"/>
<bpmn:sequenceFlow id="f2" sourceRef="t" targetRef="e" name="одобрено"/>
</bpmn:process></bpmn:definitions>"""


async def test_drawio_nodes_and_edges():
    segs = await parse("arch.drawio", DRAWIO)
    text = segs[0].text
    assert "Внутри «Agent Platform»: LangGraph" in text
    assert "- LangGraph -> LiteLLM [chat]" in text
    assert segs[0].kind == "diagram" and segs[0].locator["section"] == "Target"


async def test_bpmn_elements_and_flows():
    segs = await parse("p.bpmn", BPMN)
    text = segs[0].text
    assert "BPMN-процесс: Обработка заявки" in text
    assert "- Проверить заявку -> Готово [одобрено]" in text


async def test_xml_with_doctype_is_rejected():
    with pytest.raises(ApiError):
        await parse("x.xml", '<?xml version="1.0"?><!DOCTYPE a [<!ENTITY e "x">]><a>&e;</a>')


# ---- chunking strategy resolution ---------------------------------------------------------------
async def test_collection_default_strategy_does_not_break_other_kinds():
    segs = await parse("a.py", PY)
    for strategy in ("markdown", "sentence", "auto"):  # a markdown collection can still hold code
        chunks = chunk_segments(segs, ChunkingConfig(strategy=strategy, chunk_size=256, chunk_overlap=0))
        assert chunks
        if strategy != "sentence":
            assert all(c.kind == "code" for c in chunks)
