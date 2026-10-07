"""Excel (.xlsx) parser: sheets -> markdown tables in row groups; formulas as separate segments."""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..errors import ApiError
from .base import ParseContext, RawDocument, Segment, count_tokens

MAX_ROWS_PER_SHEET = 20_000
MAX_COLS = 60
MAX_CELL_CHARS = 500
MAX_FORMULAS = 5_000


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.6g}"
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s[: MAX_CELL_CHARS - 1] + "…" if len(s) > MAX_CELL_CHARS else s


def _md_row(cells: list[str]) -> str:
    return "| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |"


def _blocks(rows: list[tuple]) -> list[tuple[int, list[list[str]]]]:
    """Split a sheet into blocks separated by fully empty rows -> [(first_row_number, rows)]."""
    blocks: list[tuple[int, list[list[str]]]] = []
    cur: list[list[str]] = []
    first = 1
    for idx, row in enumerate(rows, start=1):
        cells = [_cell(v) for v in row[:MAX_COLS]]
        if any(cells):
            if not cur:
                first = idx
            cur.append(cells)
        elif cur:
            blocks.append((first, cur))
            cur = []
    if cur:
        blocks.append((first, cur))
    return blocks


def _record_segments(
    header: str,
    sheet: str,
    first_row: int,
    rows: list[list[str]],
    labels: list[str],
    first_col: str,
    last_col: str,
    limit: int,
) -> list[Segment]:
    """Wide/sparse sheets: `column: value` records (only non-empty cells) instead of a table full of `| |`."""
    segments: list[Segment] = []
    lines: list[str] = []
    tokens = 0
    start_row = end_row = first_row

    def flush() -> None:
        nonlocal lines, tokens
        if lines:
            loc = {"sheet": sheet, "range": f"{first_col}{start_row}:{last_col}{end_row}"}
            segments.append(Segment("\n".join(lines), kind="table", locator=loc, header=header))
        lines, tokens = [], 0

    for i, row in enumerate(rows):
        if not any(row):
            continue
        row_no = first_row + i
        if i == 0:  # may be a header or a section title: keep the raw values
            line = "- " + " | ".join(v for v in row if v)
        else:
            line = "- " + "; ".join(f"{lab}: {v}" for lab, v in zip(labels, row, strict=True) if v)
        t = count_tokens(line) + 1
        if lines and tokens + t > limit:
            flush()
        if not lines:
            start_row = row_no
        lines.append(line)
        tokens += t
        end_row = row_no
    flush()
    return segments


def _table_segments(name: str, sheet: str, first_row: int, rows: list[list[str]], limit: int) -> list[Segment]:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # spreadsheets are mostly air: drop columns that are empty in this block
    keep = [c for c in range(width) if any(r[c] for r in rows)]
    first_col, last_col = get_column_letter(keep[0] + 1), get_column_letter(keep[-1] + 1)
    rows = [[r[c] for c in keep] for r in rows]
    width = len(keep)
    head = rows[0]
    header = f"{name} / лист {sheet}"
    body_cells = [c for r in rows[1:] for c in r]
    if body_cells and (width > 10 or sum(1 for c in body_cells if not c) / len(body_cells) > 0.5):
        labels = [h[:40] or get_column_letter(keep[i] + 1) for i, h in enumerate(head)]
        return _record_segments(header, sheet, first_row, rows, labels, first_col, last_col, limit)
    prefix = f"{_md_row(head)}\n{_md_row(['---'] * width)}"
    prefix_tokens = count_tokens(prefix)
    # a heavy header (long titles) must not eat the whole chunk when repeated
    if prefix_tokens > limit * 0.4:
        short = [c[:24] for c in head]
        rep_prefix = f"{_md_row(short)}\n{_md_row(['---'] * width)}"
        if count_tokens(rep_prefix) > limit * 0.4:
            rep_prefix = ""
    else:
        rep_prefix = prefix
    rep_tokens = count_tokens(rep_prefix) if rep_prefix else 0

    segments: list[Segment] = []
    body = rows[1:] if len(rows) > 1 else []
    if not body:  # a single row: keep as a tiny table
        loc = {"sheet": sheet, "range": f"{first_col}{first_row}:{last_col}{first_row}"}
        return [Segment(prefix, kind="table", locator=loc, header=header)]

    group: list[str] = []
    group_tokens = prefix_tokens  # the first group carries the full header
    group_start = first_row + 1
    first_group = True

    def flush(end_row: int) -> None:
        nonlocal group, group_tokens, first_group
        if not group:
            return
        loc = {"sheet": sheet, "range": f"{first_col}{group_start}:{last_col}{end_row}"}
        head_text = prefix if first_group else rep_prefix
        text = (head_text + "\n" if head_text else "") + "\n".join(group)
        segments.append(Segment(text, kind="table", locator=loc, header=header))
        group, group_tokens, first_group = [], rep_tokens, False

    for offset, row in enumerate(body):
        line = _md_row(row)
        t = count_tokens(line) + 1
        row_no = first_row + 1 + offset
        if group and group_tokens + t > limit:
            flush(row_no - 1)
            group_start = row_no
        group.append(line)
        group_tokens += t
    flush(first_row + len(body))
    return segments


def _formula_segments(
    name: str, sheet: str, formulas: list[tuple[str, str, str, str]], limit: int
) -> list[Segment]:
    """formulas: (coordinate, label, formula, cached value)."""
    segments: list[Segment] = []
    group: list[tuple[str, str]] = []
    tokens = 0
    header = f"{name} / лист {sheet} / формулы"

    def flush() -> None:
        nonlocal group, tokens
        if not group:
            return
        coords = [g[0] for g in group]
        loc = {"sheet": sheet, "range": f"{coords[0]}:{coords[-1]}" if len(coords) > 1 else coords[0]}
        segments.append(Segment("\n".join(g[1] for g in group), kind="table", locator=loc, header=header))
        group, tokens = [], 0

    for coord, label, formula, value in formulas:
        line = f"{coord}" + (f" ({label})" if label else "") + f": {formula}" + (f" = {value}" if value else "")
        t = count_tokens(line) + 1
        if group and tokens + t > limit:
            flush()
        group.append((coord, line))
        tokens += t
    flush()
    return segments


def _parse_sync(raw: RawDocument, limit: int) -> list[Segment]:
    try:
        values = load_workbook(io.BytesIO(raw.data), data_only=True, read_only=False)
        formulas = load_workbook(io.BytesIO(raw.data), data_only=False, read_only=False)
    except Exception as e:  # noqa: BLE001 - bad zip, encrypted, wrong format ...
        raise ApiError(422, "validation_error", f"{raw.filename}: cannot read xlsx ({type(e).__name__})") from e

    segments: list[Segment] = []
    n_formulas = 0
    for ws in values.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows = list(ws.iter_rows(min_row=1, max_row=MAX_ROWS_PER_SHEET, values_only=True))
        for first_row, block in _blocks(rows):
            segments.extend(_table_segments(raw.filename, ws.title, first_row, block, limit))

        found: list[tuple[str, str, str, str]] = []
        for row in formulas[ws.title].iter_rows(min_row=1, max_row=MAX_ROWS_PER_SHEET):
            for c in row[:MAX_COLS]:
                if isinstance(c.value, str) and c.value.startswith("=") and n_formulas < MAX_FORMULAS:
                    n_formulas += 1
                    label = next(
                        (
                            _cell(o.value)
                            for o in row
                            if o is not c and isinstance(o.value, str) and not o.value.startswith("=") and o.value.strip()
                        ),
                        "",
                    )
                    cached = ws[c.coordinate].value
                    found.append((c.coordinate, label[:80], c.value, _cell(cached)))
        segments.extend(_formula_segments(raw.filename, ws.title, found, limit))
    values.close()
    formulas.close()
    return segments


async def parse_xlsx(raw: RawDocument, ctx: ParseContext) -> list[Segment]:
    return await asyncio.to_thread(_parse_sync, raw, max(ctx.chunk_size, 64))


EXTENSIONS = {".xlsx": parse_xlsx}
